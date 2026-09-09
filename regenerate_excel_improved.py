#!/usr/bin/env python3
"""Regenerate Excel file from local resumes using the Gemini-based pipeline."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from system import fetch_emails as pipeline


def _autosize_and_style_sheet(worksheet, df: pd.DataFrame) -> None:
    for idx, col in enumerate(df.columns, 1):
        max_length = max(df[col].astype(str).map(len).max() if not df[col].empty else 0, len(str(col)))
        col_letter = get_column_letter(idx)
        worksheet.column_dimensions[col_letter].width = min(max_length + 2, 60)

    worksheet.freeze_panes = "A2"
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    for cell in worksheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)


def regenerate_excel() -> Path:
    pipeline.require_api_key()

    if not pipeline.ATTACHMENTS_DIR.exists():
        raise RuntimeError(f"Attachments folder not found: {pipeline.ATTACHMENTS_DIR}")

    files = [p for p in pipeline.ATTACHMENTS_DIR.glob("*") if p.suffix.lower() in {".pdf", ".docx"}]
    if not files:
        raise RuntimeError(f"No resume files found in: {pipeline.ATTACHMENTS_DIR}")

    print(f"Found {len(files)} resume files. Processing...")
    records: list[dict] = []

    for file_path in files:
        print(f"- Processing: {file_path.name}")
        text = pipeline.extract_text(file_path)
        if not text.strip():
            print("  -> No text extracted, skipping")
            continue

        parsed = pipeline.parse_resume_with_gemini(text)

        records.append(
            {
                "Full Name": parsed.get("Full Name", "Unknown"),
                "Email": "",
                "Mobile Number": "",
                "Domain": parsed.get("Profession", "Unclassified"),
                "Skills": "",
                "Experience": parsed.get("Experience", ""),
                "Resume Link": str(file_path.resolve()),
                "Date Received": datetime.now().strftime("%Y-%m-%d"),
                "Profession": parsed.get("Profession", "Unclassified"),
            }
        )

    if not records:
        raise RuntimeError("No valid resumes could be processed (text extraction failed).")

    df = pd.DataFrame(records)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_file = Path("results") / "excel" / f"classified_resumes_{timestamp}.xlsx"
    output_file.parent.mkdir(parents=True, exist_ok=True)

    column_order = [
        "Full Name",
        "Email",
        "Mobile Number",
        "Domain",
        "Skills",
        "Experience",
        "Resume Link",
        "Date Received",
        "Profession",
    ]
    df = df[column_order]

    with pd.ExcelWriter(str(output_file), engine="openpyxl") as writer:
        for profession, group in df.groupby("Profession"):
            group_to_write = group.drop(columns=["Profession"]).copy()
            sheet = pipeline.safe_sheet_name(str(profession))
            group_to_write.to_excel(writer, sheet_name=sheet, index=False)
            worksheet = writer.sheets[sheet]
            _autosize_and_style_sheet(worksheet, group_to_write)

    return output_file


if __name__ == "__main__":
    print("Regenerating Excel file from local resumes (Gemini)...")
    print("-" * 60)
    try:
        output = regenerate_excel()
        print(f"\nSuccess! New Excel file created:\n{output.resolve()}")
        print("If Excel is open, close it before replacing any existing file.")
    except Exception as e:
        print(f"\nError: {e}")

